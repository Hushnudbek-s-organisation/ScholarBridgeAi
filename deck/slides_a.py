"""Slides 1-8: cover, problem, solution, journey, catalogue, trust, AI, Telegram."""

import os

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from deck_lib import (
    BODY, GOLD, GOLD_L, GREEN, INK, LIGHT, LINE, M, MUTED, NAVY, NAVY_2, SH, SW,
    TEAL, TEAL_L, WHITE, arrow, body, bullets, card, chip, footer, kicker, para,
    picture_cover, rect, rich, rule, scrim, shadow, source, stat, text, textbox, title,
)

A = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets") + "/"


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


# ------------------------------------------------------------------ 01 ------
def cover(prs):
    s = blank(prs)
    picture_cover(s, A + "hero_mit.jpg", 0, 0, SW, SH)
    scrim(s, 0, 0, SW, SH, NAVY, 76)

    # brand lock-up
    s.shapes.add_picture(A + "logo_mark.png", Inches(M), Inches(0.62),
                         height=Inches(0.52))
    text(s, M + 0.66, 0.66, 4.5, 0.44, "ScholarBridgeAI", size=17, color=WHITE,
         bold=True, font="Georgia", anchor=MSO_ANCHOR.MIDDLE)

    chip(s, 9.86, 0.72, "PITCH DECK", fill=RGBColor(0x2E, 0x44, 0x70),
         color=WHITE, size=9, w=1.46)
    chip(s, 11.44, 0.72, "2026", fill=NAVY, color=GOLD, size=9, w=0.9,
         line=RGBColor(0x4A, 0x60, 0x8C))

    text(s, M, 2.42, 9.6, 2.1,
         "Chet elda o'qish yo'li —\nendi bitta platformada.",
         size=46, color=WHITE, bold=True, font="Georgia", line_spacing=1.06)
    rect(s, M, 4.62, 1.1, 0.05, fill=GOLD)

    text(s, M, 4.92, 8.3, 1.1,
         "O'zbekiston va Markaziy Osiyo o'quvchilari uchun AI asosidagi ta'lim "
         "platformasi: universitet va grant qidiruvidan qabul, viza va jo'nashgacha.",
         size=15.5, color=RGBColor(0xD5, 0xDF, 0xEE), line_spacing=1.45)

    items = ["200 universitet · 40 mamlakat", "2 848 program", "305 grant",
             "3 til: en / uz / ru"]
    x = M
    for it in items:
        w = chip(s, x, 6.28, it, fill=RGBColor(0x2A, 0x3F, 0x68), color=WHITE,
                 size=10, h=0.4, bold=False)
        x += w + 0.13

    text(s, M, 6.92, 6.0, 0.3, "Taqdimot  ·  2026-yil oktyabr",
         size=9, color=RGBColor(0x93, 0xA4, 0xBF))
    text(s, SW - M - 3.4, 6.92, 3.4, 0.3, "scholarbridge.ai", size=9,
         color=RGBColor(0x93, 0xA4, 0xBF), align=PP_ALIGN.RIGHT)
    notes(s, "Ochilish: biz chet elda o'qish yo'lini bitta uzluksiz mahsulotga "
             "yig'ayotgan platformamiz. 30 soniya: kim uchun, nima muammo, nima "
             "qildik, qayerdaman va nima so'rayapman.")
    return s


# ------------------------------------------------------------------ 02 ------
def problem(prs):
    s = blank(prs)
    # right photo panel
    picture_cover(s, A + "campus_oxford.jpg", 8.62, 0, SW - 8.62, SH)
    scrim(s, 8.62, 0, SW - 8.62, SH, NAVY, 80)
    text(s, 9.0, 2.28, 3.7, 0.95, "100 000+", size=40, color=WHITE,
         bold=True, font="Georgia")
    text(s, 9.0, 3.2, 3.6, 0.9,
         "o'zbek talabasi bugun xorijda tahsil olmoqda",
         size=12.5, color=RGBColor(0xD5, 0xDF, 0xEE), line_spacing=1.4)
    rect(s, 9.0, 4.2, 0.62, 0.04, fill=GOLD)
    text(s, 9.0, 4.42, 3.5, 1.1,
         "Har yili 20–38 ming yangi talaba qo'shiladi.",
         size=11, color=RGBColor(0xB9, 0xC7, 0xDC), line_spacing=1.4)
    text(s, 9.0, 6.5, 3.5, 0.6, "Manba: UNESCO UIS; O'zbekiston statistika qo'mitasi (2025)",
         size=7.5, color=RGBColor(0x8F, 0x9E, 0xB6), line_spacing=1.3)

    kicker(s, M, 0.78, "Muammo")
    title(s, M, 1.2, 7.3, "Chet elda o'qish yo'li\ntarqoq va qimmat.")
    body(s, M, 2.55, 7.2,
         "O'quvchi bir vaqtning o'zida 15–20 sayt, 5 ta jadval va bir nechta "
         "maslahatchi bilan ishlashga majbur. Natija — o'tkazib yuborilgan "
         "muddatlar va noto'g'ri tanlov.", size=12.5, h=0.8)

    rows = [
        ("01", "Ma'lumot tarqoq",
         "Talablar, to'lovlar va grant muddatlari yuzlab xorijiy saytlarda — "
         "ko'pincha eskirgan."),
        ("02", "SOP to'sig'i",
         "Ingliz tilida kuchli motivatsion xat — jarayonning eng qiyin va hal "
         "qiluvchi qadami."),
        ("03", "Qimmat maslahat",
         "Agentliklar $500–2 000 so'raydi; ko'p oilalar uchun bu imkonsiz."),
        ("04", "Muddatlar o'tib ketadi",
         "Bir nechta arizani parallel yuritganda hujjat, sana va imkoniyat "
         "yo'qoladi."),
    ]
    y = 3.5
    for num, head, desc in rows:
        text(s, M, y - 0.03, 0.5, 0.3, num, size=10.5, color=TEAL, bold=True,
             spacing=60)
        text(s, M + 0.62, y - 0.08, 6.6, 0.3, head, size=13.5, color=INK, bold=True)
        text(s, M + 0.62, y + 0.24, 6.62, 0.5, desc, size=10.5, color=BODY,
             line_spacing=1.35)
        rule(s, M, y + 0.78, 7.2)
        y += 0.85

    footer(s, 2)
    notes(s, "Muammo 4 qismga bo'linadi. Har biriga mahsulotda aniq javob bor. "
             "Muhim: bu bozor 'yangi' emas — u shunchaki hali dasturiy ta'minot "
             "bilan hal qilinmagan.")
    return s


# ------------------------------------------------------------------ 03 ------
def solution(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, 0, 0.16, SH, fill=TEAL)

    kicker(s, M, 0.72, "Yechim", color=GOLD)
    title(s, M, 1.14, 8.6, "Bitta profil. Butun sayohat.", color=WHITE)
    text(s, M, 2.02, 9.4, 0.8,
         "Bitta hisob — birinchi qidiruvdan jo'nash kunigacha. Har bir tavsiya "
         "o'quvchining haqiqiy ma'lumotiga tayanadi, umumiy maslahatga emas.",
         size=13, color=RGBColor(0xC9, 0xD5, 0xE6), line_spacing=1.45)

    feats = [
        ("01", "Moslik dvigateli",
         "GPA, test, budjet va til talablari bo'yicha 4 alohida o'lchov — "
         "qabul ehtimoli emas."),
        ("02", "Grant katalogi",
         "305 grant: qamrov, talablar, muddat va rasmiy manba bilan."),
        ("03", "AI studiya",
         "SOP drafti, insho tekshiruvi, profil baholash va rubrika bo'yicha fikr."),
        ("04", "Ariza boshqaruvi",
         "Har bir ariza uchun talablar, hujjatlar, topshirish va natija."),
        ("05", "Viza tayyorlovi",
         "Ovozli konsullik intervyusi simulyatori va mamlakat talablari."),
        ("06", "Telegram",
         "Bot va Mini App: eslatmalar, qidiruv va AI maslahatchi — bir xil hisob."),
    ]
    x0, y0 = M, 3.02
    cw, ch, gx, gy = 3.86, 1.58, 0.22, 0.18
    for i, (num, head, desc) in enumerate(feats):
        x = x0 + (i % 3) * (cw + gx)
        y = y0 + (i // 3) * (ch + gy)
        card(s, x, y, cw, ch, fill=NAVY_2, line=RGBColor(0x33, 0x4A, 0x74),
             radius=0.07, lift=False)
        rect(s, x, y, 0.05, ch, fill=GOLD if i % 2 == 0 else TEAL)
        text(s, x + 0.3, y + 0.2, 0.6, 0.24, num, size=9.5, color=GOLD, bold=True,
             spacing=80)
        text(s, x + 0.3, y + 0.46, cw - 0.55, 0.3, head, size=13.5, color=WHITE,
             bold=True)
        text(s, x + 0.3, y + 0.8, cw - 0.55, 0.72, desc, size=9.8,
             color=RGBColor(0xB9, 0xC7, 0xDC), line_spacing=1.34)

    text(s, SW - M - 4.6, 6.86, 4.6, 0.3,
         "Bugun jonli: 14 bo'lim · 3 til · 98 jadval · Payme/Click",
         size=8.5, color=RGBColor(0x8F, 0x9E, 0xB6), align=PP_ALIGN.RIGHT)
    notes(s, "Har bir karta — saytda ishlaydigan alohida bo'lim. Muhim farq: "
             "moslik 'ehtimol' emas, o'lchov; ma'lumot esa manba bilan.")
    return s


# ------------------------------------------------------------------ 04 ------
def journey(prs):
    s = blank(prs)
    kicker(s, M, 0.78, "Mahsulot")
    title(s, M, 1.2, 11.2, "Sakkiz bosqich — bitta uzluksiz yo'l")
    text(s, M, 2.0, 11.4, 0.5,
         "Platforma o'quvchining o'z ma'lumotidan qaysi bosqichda turganini "
         "avtomatik hisoblaydi — hech narsa qo'lda belgilanmaydi.",
         size=12.5, color=BODY, line_spacing=1.4)

    stages = [
        ("01", "Kashf etish", "200 universitet va 2 848 program: davlat, daraja, "
                              "yo'nalish va to'lov bo'yicha filtr."),
        ("02", "Moslashtirish", "Profilga moslik 4 o'lchovda; har bir natijada "
                                "manba va tekshirilgan sana."),
        ("03", "Tayyorlanish", "10 fazali o'quv reja, test rejasi, hujjatlar "
                               "ro'yxati va keyingi qadam."),
        ("04", "Ariza", "Har bir ariza uchun alohida ish maydoni: talablar, "
                        "hujjatlar, insholar, topshirish."),
        ("05", "Qabul", "Takliflar, shartlar va qarorlar bitta joyda."),
        ("06", "Moliyalashtirish", "To'lov rejasi, grant qamrovi va moliyaviy "
                                   "bo'shliq hisobi."),
        ("07", "Viza", "Mamlakat talablari, hujjat matritsasi va intervyu "
                       "simulyatori."),
        ("08", "Jo'nash", "Parvoz, turar joy, sug'urta va 7/30/90 kunlik "
                          "kelish ro'yxati."),
    ]
    x0, y0 = M, 2.72
    cw, ch, gx, gy = 2.9, 1.62, 0.2, 0.22
    for i, (num, head, desc) in enumerate(stages):
        row, col = divmod(i, 4)
        x = x0 + col * (cw + gx)
        y = y0 + row * (ch + gy)
        accent = TEAL if row == 0 else GREEN
        card(s, x, y, cw, ch, fill=WHITE, line=LINE, radius=0.07)
        rect(s, x, y, cw, 0.05, fill=accent)
        text(s, x + 0.24, y + 0.24, 0.6, 0.22, num, size=9, color=accent,
             bold=True, spacing=80)
        text(s, x + 0.24, y + 0.5, cw - 0.44, 0.3, head, size=12.5, color=INK,
             bold=True)
        text(s, x + 0.24, y + 0.84, cw - 0.44, 0.7, desc, size=9.2, color=BODY,
             line_spacing=1.32)
        if col < 3:
            arrow(s, x + cw + 0.035, y + ch / 2 - 0.07, 0.13, 0.14,
                  color=RGBColor(0xC9, 0xD3, 0xE0))

    rect(s, M, 6.28, SW - 2 * M, 0.72, fill=LIGHT, radius=0.28)
    text(s, M + 0.32, 6.28, SW - 2 * M - 0.64, 0.72,
         "Bir xil ma'lumot hamma joyda: bir marta kiritilgan profil butun "
         "yo'l davomida qayta ishlatiladi.",
         size=11, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 4)
    notes(s, "Bu 8 bosqich — saytdagi navigatsiya guruhlari bilan bir xil. "
             "Ko'rsatish: bitta o'quvchi profilini to'ldirgach, bosqich o'zi "
             "siljiydi.")
    return s


# ------------------------------------------------------------------ 05 ------
def catalogue(prs):
    s = blank(prs)
    kicker(s, M, 0.78, "Mahsulot · Katalog")
    title(s, M, 1.2, 11.0, "Ma'lumot — to'qilgan emas, manbali")
    text(s, M, 2.0, 11.2, 0.5,
         "Har bir yozuv rasmiy manbaga bog'langan va oxirgi tekshirilgan sana "
         "bilan ko'rsatiladi. Eskirgan ma'lumot «tekshirilmagan» deb belgilanadi.",
         size=12, color=BODY, line_spacing=1.4)

    stats = [("200", "universitet", TEAL), ("2 848", "program", TEAL),
             ("305", "grant va stipendiya", GREEN), ("40", "mamlakat", NAVY)]
    for i, (v, l, c) in enumerate(stats):
        x = M + i * 2.98
        card(s, x, 2.72, 2.78, 1.18, fill=WHITE, line=LINE, radius=0.08)
        text(s, x + 0.28, 2.86, 2.3, 0.5, v, size=30, color=c, bold=True,
             font="Georgia")
        text(s, x + 0.28, 3.42, 2.3, 0.3, l, size=10, color=MUTED)

    photos = [
        ("campus_cambridge.jpg", "Kembrij universiteti", "Buyuk Britaniya", "QS #6"),
        ("campus_oxford.jpg", "Oksford universiteti", "Buyuk Britaniya", "QS #4"),
        ("campus_eth.jpg", "ETH Zürich", "Shveytsariya", "QS #7"),
    ]
    for i, (f, name, country, rank) in enumerate(photos):
        x = M + i * 4.0
        picture_cover(s, A + f, x, 4.08, 3.82, 1.98)
        chip(s, x + 0.16, 4.22, rank, fill=WHITE, color=NAVY, size=8.5, h=0.26,
             w=0.72)
        text(s, x, 6.16, 3.82, 0.28, name, size=10.5, color=INK, bold=True)
        text(s, x, 6.42, 3.82, 0.24, country, size=9, color=MUTED)

    source(s, M, 6.78, 9.4,
           "Foto: Wikimedia Commons — sayt katalogidagi rasmlar. Ma'lumot: platforma katalogi, 2026-10.")
    footer(s, 5)
    notes(s, "Rasmlar — saytdagi kampus fotolari (Wikimedia Commons). "
             "Raqamlar jonli bazadan: 200 universitet, 2 848 program, 305 grant, "
             "40 mamlakat.")
    return s


# ------------------------------------------------------------------ 06 ------
def trust(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    picture_cover(s, A + "hero_mit.jpg", 9.98, 0, SW - 9.98, SH)
    scrim(s, 9.98, 0, SW - 9.98, SH, NAVY, 72)
    text(s, 10.34, 5.6, 2.5, 0.9,
         "MIT — katalogdagi 200\nuniversitetdan biri.",
         size=8.5, color=RGBColor(0xB9, 0xC7, 0xDC), line_spacing=1.35)

    kicker(s, M, 0.78, "Farq", color=GOLD)
    title(s, M, 1.2, 8.0, "Ishonch — mahsulotning o'zi.", color=WHITE)
    text(s, M, 2.08, 7.4, 0.7,
         "Ta'lim bozorida eng qimmat narsa — noto'g'ri ma'lumot. Biz to'rt "
         "qoidani kodda majburlaymiz, marketingda emas.",
         size=12.5, color=RGBColor(0xC9, 0xD5, 0xE6), line_spacing=1.42)

    cols = [
        ("01", "Har faktda manba",
         "Har bir talab yonida rasmiy havola va oxirgi tekshirilgan sana. "
         "Manbasiz yozuv «ko'rsatilmagan» bo'lib qoladi."),
        ("02", "O'ylab topilgan ehtimol yo'q",
         "Qabul ehtimoli foizi ko'rsatilmaydi: tasdiqlangan metodologiya "
         "bo'lmaguncha u o'chirilgan."),
        ("03", "AI javobi yorliqlanadi",
         "verified_fact · user_provided · ai_suggestion · uncertain — har bir "
         "gap qayerdan kelgani ko'rinadi."),
        ("04", "Maxfiylik",
         "Har bir so'rov egasi serverda tekshiriladi; AI kalitlari AES-256 "
         "bilan shifrlangan va hech qachon javobda chiqmaydi."),
    ]
    x = M
    for i, (num, head, desc) in enumerate(cols):
        text(s, x, 3.16, 2.0, 0.24, num, size=9.5, color=GOLD, bold=True, spacing=80)
        text(s, x, 3.44, 2.12, 0.62, head, size=11.5, color=WHITE, bold=True,
             line_spacing=1.2)
        text(s, x, 4.12, 2.1, 1.5, desc, size=8.7, color=RGBColor(0xB9, 0xC7, 0xDC),
             line_spacing=1.36)
        if i < 3:
            rect(s, x + 2.22, 3.2, 0.012, 2.4, fill=RGBColor(0x3A, 0x4F, 0x78))
        x += 2.41

    rect(s, M, 5.82, 8.4, 0.92, fill=RGBColor(0x24, 0x38, 0x60), radius=0.12)
    rect(s, M, 5.82, 0.05, 0.92, fill=GOLD)
    text(s, M + 0.34, 5.82, 7.8, 0.56,
         "«Bilmagan narsamizni to'qimaymiz.»", size=15, color=WHITE,
         bold=True, font="Georgia", anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + 0.34, 6.34, 7.8, 0.3, "— mahsulot qoidasi, avtomatik testlar bilan",
         size=9, color=RGBColor(0x8F, 0x9E, 0xB6))

    footer(s, 6, dark=True)
    notes(s, "Bu slayd investor uchun eng muhim farq: raqobatchilar 'AI "
             "bashorat' bilan sotadi, biz esa faqat tekshiriladigan narsani "
             "ko'rsatamiz. Bu uzoq muddatda ishonchni ushlab turadi.")
    return s


# ------------------------------------------------------------------ 07 ------
def ai(prs):
    s = blank(prs)
    kicker(s, M, 0.78, "Arxitektura")
    title(s, M, 1.2, 6.4, "AI — mahsulot ichida,\nlekin nazorat ostida")
    text(s, M, 2.3, 5.9, 0.8,
         "Bitta umumiy «chat» emas: har bir vazifa o'z provayderiga, o'z "
         "byudjeti va o'z yorlig'i bilan boradi.",
         size=11.5, color=BODY, line_spacing=1.4)

    # diagram
    bx, bw = M, 5.86
    boxes = [
        ("Foydalanuvchi", "sayt · Telegram bot · Mini App", LIGHT, INK),
        ("Next.js API", "sessiya + profil egaligi + kunlik kvota", TEAL_L, TEAL),
        ("AI router", "vazifaga qarab provayder tanlaydi", RGBColor(0xE9, 0xEE, 0xF7), NAVY),
    ]
    y = 3.28
    for i, (head, sub, fill, col) in enumerate(boxes):
        shp = rect(s, bx, y, bw, 0.72, fill=fill, line=LINE, radius=0.1)
        text(s, bx + 0.26, y + 0.1, bw - 0.5, 0.28, head, size=11.5, color=col,
             bold=True)
        text(s, bx + 0.26, y + 0.38, bw - 0.5, 0.26, sub, size=8.8, color=BODY)
        if i < 2:
            arrow(s, bx + bw / 2 - 0.08, y + 0.76, 0.16, 0.14,
                  color=RGBColor(0xBF, 0xCB, 0xDA))
        y += 0.9

    provs = [("Groq", 0.86), ("OpenRouter", 1.32), ("OpenAI", 1.0),
             ("Anthropic", 1.24), ("Gemini", 0.96)]
    px = bx
    for p, pw in provs:
        chip(s, px, y + 0.06, p, fill=WHITE, color=NAVY, size=8.6, h=0.32,
             w=pw, line=LINE, bold=False)
        px += pw + 0.1

    # right column
    rx = 7.62
    rect(s, rx, 1.3, 0.012, 5.2, fill=LINE)
    items = [
        ("Vazifaga alohida provayder",
         "admissions, essay, document, general — admin panelidan deploy'siz "
         "almashtiriladi."),
        ("Kunlik kvota va byudjet",
         "bepul: 5 so'rov/kun · Premium: 60 so'rov/kun. Xarajat so'rov "
         "boshiga yoziladi."),
        ("Aniq failover",
         "zaxira faqat sozlanganda ishlaydi va javob «provider:fallback» deb "
         "belgilanadi."),
        ("Kalitlar serverda",
         "AES-256-GCM; javobda, logda va xatoda ko'rinmaydi. Har bir so'rov "
         "audit qilinadi."),
        ("AI javob yorlig'i",
         "model taklifi fakt bilan aralashmaydi; foydalanuvchi nima AI fikri "
         "ekanini ko'radi."),
    ]
    y2 = 1.3
    for head, desc in items:
        text(s, rx + 0.32, y2, 4.6, 0.28, head, size=11.5, color=INK, bold=True)
        text(s, rx + 0.32, y2 + 0.28, 4.6, 0.6, desc, size=9.4, color=BODY,
             line_spacing=1.36)
        y2 += 1.02

    footer(s, 7)
    notes(s, "Texnik investor uchun: provider-agnostic router, kvota, xarajat "
             "hisobi va shifrlangan kalitlar. Hozir production'da Groq ishlaydi, "
             "lekin almashtirish bir tugma.")
    return s


# ------------------------------------------------------------------ 08 ------
def telegram(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=LIGHT)
    kicker(s, M, 0.78, "Ikkinchi eshik")
    title(s, M, 1.2, 7.0, "Telegram — asosiy kanal")
    text(s, M, 2.06, 6.6, 0.9,
         "O'zbekistonda o'quvchi Instagram va Telegram'da yashaydi. Bot va Mini "
         "App saytning o'z API'sini chaqiradi: bir xil hisob, bir xil qidiruv, "
         "bir xil kvota — o'quvchi yangi ilova o'rnatmaydi.",
         size=11.5, color=BODY, line_spacing=1.45)

    items = [
        "Eslatmalar: muddatdan 30/14/7/3/1/0 kun oldin, bitta xabar — bir marta.",
        "Qidiruv va saqlanganlar bot ichida: /universities, /scholarships, /next.",
        "AI maslahatchi — saytdagi bilan bir xil kvota va qoidalar.",
        "6 xonali kod bilan saytga kirish; har bir ulanish audit qilinadi.",
    ]
    bullets(s, M, 3.3, 6.6, items, size=10.5, gap=0.16, h=3.0)

    chip(s, M, 5.9, "1 hisob = 1 Telegram", fill=WHITE, color=NAVY, size=9.5,
         h=0.36, line=LINE)
    chip(s, M + 2.4, 5.9, "Mini App: /tg", fill=WHITE, color=NAVY, size=9.5,
         h=0.36, line=LINE)

    # phone mock
    px, py, pw, ph = 7.86, 0.82, 3.5, 5.86
    rect(s, px, py, pw, ph, fill=NAVY, radius=0.08)
    rect(s, px + 0.11, py + 0.11, pw - 0.22, ph - 0.22, fill=WHITE, radius=0.05)
    rect(s, px, py, pw, 0.6, fill=NAVY, radius=0.08)
    rect(s, px, py + 0.42, pw, 0.18, fill=NAVY)
    text(s, px + 0.22, py + 0.09, 2.6, 0.26, "ScholarBridge", size=10.5,
         color=WHITE, bold=True)
    text(s, px + 0.22, py + 0.33, 2.6, 0.22, "bot · @scholarbridge_bot", size=7.4,
         color=RGBColor(0xA9, 0xB8, 0xD2))

    bubbles = [
        ("in", "MITda kompyuter fanlari bo'yicha\ngrant bormi?", 0.86, 0.62),
        ("out", "3 ta mos grant topdim.\nEng yaqin muddat — 12-noyabr.\nEslatma qo'shamanmi?", 1.62, 0.82),
        ("in", "Ha, iltimos", 2.58, 0.44),
        ("out", "✓ Eslatma qo'shildi: 30/14/7/3/1 kun\noldin xabar beraman.", 3.16, 0.7),
    ]
    for kind, txt, dy, bh in bubbles:
        yy = py + dy
        if kind == "in":
            rect(s, px + 1.12, yy, 2.24, bh, fill=RGBColor(0xE8, 0xEE, 0xF6),
                 radius=0.18)
            text(s, px + 1.28, yy + 0.08, 1.96, bh - 0.12, txt, size=7.4,
                 color=INK, line_spacing=1.3)
        else:
            rect(s, px + 0.2, yy, 2.32, bh, fill=RGBColor(0xDC, 0xEE, 0xF5),
                 radius=0.16)
            text(s, px + 0.34, yy + 0.08, 2.06, bh - 0.12, txt, size=7.4,
                 color=INK, line_spacing=1.3)

    rect(s, px + 0.2, py + 4.16, pw - 0.4, 1.34, fill=WHITE, line=LINE, radius=0.1)
    text(s, px + 0.34, py + 4.26, pw - 0.7, 0.24, "BUYRUQLAR", size=7.4,
         color=MUTED, bold=True, spacing=120)
    text(s, px + 0.34, py + 4.54, pw - 0.7, 0.8,
         "/next  ·  /deadlines  ·  /saved\n/advisor  ·  /applications  ·  /help",
         size=7.8, color=TEAL, line_spacing=1.5)

    text(s, M, 6.42, 6.6, 0.26,
         "Mini App faqat web.telegram.org tomonidan ramkaga olinadi.",
         size=8, color=MUTED)
    footer(s, 8)
    notes(s, "O'zbekistonda Telegram — asosiy o'quv kanali. Shuning uchun bot "
             "mahsulotning 'ikkinchi eshigi', alohida mahsulot emas: kod ham, "
             "qoidalar ham bir xil.")
    return s
