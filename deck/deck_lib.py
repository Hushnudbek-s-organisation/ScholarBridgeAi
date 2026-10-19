"""Design system + drawing helpers for the ScholarBridgeAI pitch deck.

Everything is drawn with native PowerPoint shapes / text boxes / tables /
charts so the .pptx stays fully editable (and imports cleanly into Google
Slides). No slide is a flat image.
"""

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from PIL import Image

# ----------------------------------------------------------------- brand ----
NAVY = RGBColor(0x1B, 0x2A, 0x4A)        # logo background
NAVY_2 = RGBColor(0x22, 0x36, 0x5E)      # lighter navy for panels
INK = RGBColor(0x0F, 0x17, 0x2A)
BODY = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x7C, 0x8A, 0x9E)
LINE = RGBColor(0xDD, 0xE4, 0xEC)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x1C, 0x72, 0x93)
TEAL_L = RGBColor(0xE8, 0xF2, 0xF7)
GOLD = RGBColor(0xE8, 0xB9, 0x4E)
GOLD_L = RGBColor(0xFB, 0xF2, 0xDE)
GREEN = RGBColor(0x1B, 0x8A, 0x6B)
GREEN_L = RGBColor(0xE7, 0xF5, 0xF0)

FONT = "Arial"
FONT_HEAD = "Georgia"     # headings — ships with Office and Google Slides

SW, SH = 13.333, 7.5      # slide size, inches
M = 0.72                  # page margin

SLIDE_W = Inches(SW)
SLIDE_H = Inches(SH)


# ------------------------------------------------------------- primitives ---
def _alpha(shape, pct):
    """Apply alpha (%) to every solid colour in the shape's spPr."""
    spPr = shape._element.spPr
    val = str(int(round(pct * 1000)))
    for srgb in spPr.iter(qn("a:srgbClr")):
        for old in srgb.findall(qn("a:alpha")):
            srgb.remove(old)
        a = etree.SubElement(srgb, qn("a:alpha"))
        a.set("val", val)


def shadow(shape, blur=14, dist=4, alpha=10, direction=5400000):
    """Soft outer shadow (Office native), removes the inherited one first."""
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    shd = etree.SubElement(eff, qn("a:outerShdw"))
    shd.set("blurRad", str(int(Pt(blur).emu)))
    shd.set("dist", str(int(Pt(dist).emu)))
    shd.set("dir", str(direction))
    shd.set("algn", "ctr")
    shd.set("rotWithShape", "0")
    clr = etree.SubElement(shd, qn("a:srgbClr"))
    clr.set("val", "1B2A4A")
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 1000)))


def no_shadow(shape):
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    etree.SubElement(spPr, qn("a:effectLst"))


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, radius=None,
         alpha=None, shape_type=None, dash=None):
    kind = shape_type or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if alpha is not None:
            _alpha(shp, alpha)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
        if dash is not None:
            shp.line.dash_style = dash
    no_shadow(shp)
    shp.text_frame.word_wrap = True
    return shp


def card(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.06, lift=True):
    shp = rect(slide, x, y, w, h, fill=fill, line=line, lw=1, radius=radius)
    if lift:
        shadow(shp, blur=16, dist=4, alpha=8)
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0)):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    return tb


def _apply_run(run, size, color, bold=False, font=FONT, italic=False, spacing=None):
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    if spacing is not None:
        run.font._rPr.set("spc", str(int(spacing)))


def para(tf, first=False, align=PP_ALIGN.LEFT, space_before=0, space_after=0,
         line_spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line_spacing:
        p.line_spacing = line_spacing
    return p


def rich(tf, chunks, first=False, align=PP_ALIGN.LEFT, space_before=0,
         space_after=0, line_spacing=None):
    """chunks = [(text, {size,color,bold,font,italic,spacing}), ...]"""
    p = para(tf, first=first, align=align, space_before=space_before,
             space_after=space_after, line_spacing=line_spacing)
    for text, style in chunks:
        r = p.add_run()
        r.text = text
        _apply_run(
            r,
            style.get("size", 14),
            style.get("color", BODY),
            style.get("bold", False),
            style.get("font", FONT),
            style.get("italic", False),
            style.get("spacing"),
        )
    return p


def text(slide, x, y, w, h, value, size=14, color=BODY, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15,
         spacing=None, italic=False, margins=(0, 0, 0, 0)):
    tb = textbox(slide, x, y, w, h, anchor=anchor, margins=margins)
    lines = value.split("\n")
    for i, line in enumerate(lines):
        rich(tb.text_frame, [(line, dict(size=size, color=color, bold=bold,
                                        font=font, italic=italic, spacing=spacing))],
             first=(i == 0), align=align, line_spacing=line_spacing)
    return tb


# --------------------------------------------------------------- blocks -----
def kicker(slide, x, y, label, color=TEAL, w=6.0):
    """Small uppercase eyebrow with a short rule."""
    text(slide, x, y, w, 0.26, label.upper(), size=10, color=color, bold=True,
         spacing=190, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, x, y + 0.34, 0.52, 0.045, fill=color)


def _wrapped_lines(value, size, w, factor=0.0122):
    """Rough line count for a bold serif heading at `size` pt in `w` inches."""
    per_line = max(6, int(w / (size * factor)))
    total = 0
    for seg in value.split("\n"):
        total += max(1, -(-len(seg) // per_line))
    return total


def title(slide, x, y, w, value, color=INK, size=31, font=FONT_HEAD):
    lines = _wrapped_lines(value, size, w)
    return text(slide, x, y, w, 0.60 * lines + 0.12, value, size=size,
                color=color, bold=True, font=font, line_spacing=1.08)


def title_h(value, w, size=31):
    """Height a title will need — lets callers place the next element."""
    return 0.60 * _wrapped_lines(value, size, w) + 0.12


def body(slide, x, y, w, value, color=BODY, size=13.5, h=None, line_spacing=1.4):
    lines = int(len(value) / (w * 12)) + value.count("\n") + 2
    return text(slide, x, y, w, h or 0.26 * lines, value, size=size,
                color=color, line_spacing=line_spacing)


def bullets(slide, x, y, w, items, size=13, color=BODY, gap=0.12,
            dot=TEAL, dot_size=0.075, line_spacing=1.32, h=4.0):
    """Custom bullets (circle + text) so indenting behaves predictably."""
    cy = y
    for item in items:
        rect(slide, x, cy + 0.112, dot_size, dot_size, fill=dot, shape_type=MSO_SHAPE.OVAL)
        tb = textbox(slide, x + dot_size + 0.13, cy - 0.02, w - dot_size - 0.13, 0.3)
        rich(tb.text_frame, [(item, dict(size=size, color=color))], first=True,
             line_spacing=line_spacing)
        est_lines = max(1, int(len(item) / max(12, (w - 0.3) * 11.5)) + 1)
        cy += 0.245 * est_lines + gap
    return cy


def chip(slide, x, y, label, fill=TEAL_L, color=TEAL, size=10.5, pad=0.12,
         h=0.32, w=None, bold=True, radius=0.5, line=None):
    width = w if w else (len(label) * (size * 0.0112) + pad * 2 + 0.08)
    rect(slide, x, y, width, h, fill=fill, line=line, radius=radius)
    text(slide, x, y, width, h, label, size=size, color=color, bold=bold,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return width


def stat(slide, x, y, w, value, label, color=TEAL, value_size=34, label_size=10.5,
         label_color=MUTED):
    text(slide, x, y, w, 0.52, value, size=value_size, color=color, bold=True,
         font=FONT_HEAD)
    text(slide, x, y + 0.56, w, 0.5, label, size=label_size, color=label_color,
         line_spacing=1.2)


def picture_cover(slide, path, x, y, w, h, border=None):
    """Insert an image cropped to exactly fill the box (like CSS cover)."""
    iw, ih = Image.open(path).size
    box_ratio = w / h
    img_ratio = iw / ih
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if img_ratio > box_ratio:          # image too wide -> trim sides
        crop = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    elif img_ratio < box_ratio:        # image too tall -> trim top/bottom
        crop = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(1)
    return pic


def scrim(slide, x, y, w, h, color=NAVY, alpha=72):
    s = rect(slide, x, y, w, h, fill=color, alpha=alpha)
    slide.shapes._spTree.remove(s._element)
    slide.shapes._spTree.append(s._element)   # keep on top of pictures
    return s


def footer(slide, page, label="ScholarBridgeAI  ·  Pitch deck 2026", dark=False):
    col = RGBColor(0x8F, 0x9E, 0xB6) if dark else MUTED
    line_col = RGBColor(0x33, 0x45, 0x68) if dark else LINE
    rect(slide, M, SH - 0.62, SW - 2 * M, 0.012, fill=line_col)
    text(slide, M, SH - 0.52, 7.5, 0.26, label, size=8.5, color=col)
    text(slide, SW - M - 1.2, SH - 0.52, 1.2, 0.26, f"{page:02d}", size=8.5,
         color=col, align=PP_ALIGN.RIGHT)


def source(slide, x, y, w, value, dark=False):
    text(slide, x, y, w, 0.5, value, size=8.5,
         color=RGBColor(0x8F, 0x9E, 0xB6) if dark else MUTED, line_spacing=1.25)


def rule(slide, x, y, w, color=LINE, h=0.012):
    rect(slide, x, y, w, h, fill=color)


def arrow(slide, x, y, w, h, color=RGBColor(0xC3, 0xCE, 0xDB)):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    no_shadow(shp)
    return shp
